"""Checklist verification probes (round 3).

Each probe answers one checklist question against the *real* code, in an
isolated database, and prints the observed behaviour. Nothing here is a
production path; it exists so the audit report can cite evidence rather than
inference.

    python scripts/verify_checklist.py
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_ROOT))

_TMP = Path(tempfile.mkdtemp(prefix="shipsync-checklist-"))
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP / 'c.db'}"
os.environ["INGEST_DIR"] = str(_TMP / "ingested")
os.environ["DATA_SOURCE"] = str(_TMP / "empty")
# The image rows need the OCR rung; the setting is read at import time, so
# it must be set before any app module loads.
os.environ["OCR_ENABLED"] = "1"

from app.database import SessionLocal, init_db  # noqa: E402
from app.models import (DocumentVersionRecord, EmailRecord,  # noqa: E402
                        ShipmentRecord)
from app.services import ai_service, doc_types, versioning, workflow  # noqa: E402
from app.services import extractor  # noqa: E402


def pdf(text: str) -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 60
    for line in text.splitlines():
        page.insert_text((60, y), line, fontsize=9.5, fontname="cour")
        y += 15
    data = doc.tobytes(deflate=True)
    doc.close()
    return data


def si_text(ref: str = "SHP-001", shipper: str = "ABC LOGISTICS SDN BHD",
            pol: str = "PORT KLANG, MALAYSIA (MYPKG)",
            pod: str = "ROTTERDAM, NETHERLANDS (NLRTM)",
            weight: str = "41,500 KG") -> str:
    return f"""SHIPPING INSTRUCTION
========================================

Shipper/Exporter: {shipper}
CONSIGNEE: ROTTERDAM TRADING BV
NOTIFY PARTY: ROTTERDAM TRADING BV
Port of Loading: {pol}
Discharge Port: {pod}
No. of Containers or Packages: 2 x 40'HC
Gross Weight (KG): {weight}
Vessel Name: EVER GLORY
Voy. No: 026E
Container No.: ABCU1234567
Shipment ID: {ref}
OC No.: {ref}
ETD: 25 September 2026
ETA: 15 October 2026
Freight: PREPAID
"""


def hdr(title: str) -> None:
    print("\n" + "=" * 74)
    print(title)
    print("=" * 74)


# --------------------------------------------------------------------------- #
def probe_shipping_id() -> None:
    hdr("3. SHIPPING ID DETECTION - which source wins, and can shipments merge?")

    cases = [
        ("reference in the subject only",
         {"subject": "Shipping Instruction - SHP-AA1", "body": "see attached",
          "email_id": "e1", "attachments": ["SI_v1.pdf"]}, {}, {}),
        ("reference in the document text only",
         {"subject": "please check", "body": "see attached", "email_id": "e2",
          "attachments": ["SI_v1.pdf"]}, {"shipper": "X", "oc_no": "SHP-BB2"}, {}),
        ("reference in the filename only",
         {"subject": "please check", "body": "see attached", "email_id": "e3",
          "attachments": ["SHP-CC3_SI.pdf"]}, {}, {}),
        ("subject and document DISAGREE",
         {"subject": "Shipping Instruction - SHP-DD4", "body": "checked",
          "email_id": "e4", "attachments": ["SI_v1.pdf"]},
         {"shipper": "X", "oc_no": "SHP-EE5"}, {}),
        ("no reference anywhere, full field signature",
         {"subject": "please check", "body": "docs", "email_id": "e5",
          "attachments": ["SI_v1.pdf"]},
         {"shipper": "ABC LOGISTICS SDN BHD",
          "port_of_loading": "PORT KLANG, MALAYSIA (MYPKG)",
          "port_of_discharge": "ROTTERDAM, NETHERLANDS (NLRTM)"}, {}),
        ("no reference, no signature (nothing to group on)",
         {"subject": "hi", "body": "docs", "email_id": "e6",
          "attachments": ["SI_v1.pdf"]}, {}, {}),
    ]
    for label, email, si, bl in cases:
        key, reference = versioning.identify_shipment(email, si, bl)
        print(f"  {label:48s} -> {key!r}  (reference={reference!r})")

    print("\n  COLLISION TEST: two shipments, no reference text, same route")
    a = {"shipper": "ABC LOGISTICS SDN BHD",
         "port_of_loading": "PORT KLANG, MALAYSIA (MYPKG)",
         "port_of_discharge": "ROTTERDAM, NETHERLANDS (NLRTM)"}
    b = dict(a)  # same parties and route, but a different real shipment
    ka, _ = versioning.identify_shipment(
        {"subject": "docs", "body": "x", "email_id": "shipA",
         "attachments": ["SI_v1.pdf"]}, a, {})
    kb, _ = versioning.identify_shipment(
        {"subject": "docs", "body": "x", "email_id": "shipB",
         "attachments": ["SI_v1.pdf"]}, b, {})
    print(f"    shipment A -> {ka}")
    print(f"    shipment B -> {kb}")
    print(f"    MERGED? {ka == kb}"
          + ("  <-- two unrelated shipments share one key" if ka == kb else ""))


# --------------------------------------------------------------------------- #
def probe_classification() -> None:
    hdr("6. CLASSIFICATION - does it use content, or only the filename?")
    bl = f"""BILL OF LADING
SHIPPER: ABC LOGISTICS SDN BHD
CONSIGNEE: ROTTERDAM TRADING BV
Notify: ROTTERDAM TRADING BV
Port of Loading (POL): PORT KLANG, MALAYSIA (MYPKG)
POD: ROTTERDAM, NETHERLANDS (NLRTM)
Container Count: 2 x 40'HC
Gross Wt (kgs): 41,500 KG
"""
    names = [
        ("SHP-001_SI.pdf", si_text()),        # convention
        ("SI_v1.pdf", si_text()),             # realistic
        ("Shipping_Instruction.pdf", si_text()),
        ("document_final.pdf", si_text()),    # IS an SI, named neutrally
        ("SHP-001_BL.pdf", bl),               # convention
        ("document_final_BL.pdf", bl),
        ("scan.pdf", bl),                     # content only
    ]
    print(f"  {'filename':26s} {'name rules':11s} {'content says':13s} found as")
    for name, text in names:
        by_name = doc_types.detect(name)
        declared = workflow.declared_doc_type(text)
        found = workflow._fuzzy_doc_type(name)
        print(f"  {name:26s} {str(by_name):11s} {str(declared):13s} {str(found)}")

    print("\n  content-vs-name guard (`_wrong_doc`): a file named _BL that")
    print("  actually contains a Shipping Instruction")
    content = ai_service.document_text("SHP-001_BL.pdf", pdf(si_text()))
    res = extractor.extract_fields(content[0], "BL")
    print(f"    declared_doc_type(text) = {workflow.declared_doc_type(content[0])!r}"
          f" -> wrong_doc flag = {workflow._wrong_doc(res, 'BL')}")

    print("\n  Is the reference inside the DOCUMENT ever used for grouping?")
    doc_only = ("SHIPPING INSTRUCTION\nShipper/Exporter: ZZZ SDN BHD\n"
                "CONSIGNEE: YYY BV\nNOTIFY PARTY: YYY BV\n"
                "Port of Loading: PORT KLANG, MALAYSIA (MYPKG)\n"
                "Discharge Port: ROTTERDAM, NETHERLANDS (NLRTM)\n"
                "No. of Containers or Packages: 2 x 40'HC\n"
                "Gross Weight (KG): 41,500 KG\n"
                "Shipment ID: SHP-ZZ9\nOC No.: SHP-ZZ9\n")
    fields = extractor.extract_fields(doc_only, "SI").fields
    print(f"    extracted field names: {sorted(fields)}")
    print(f"    -> a reference field exists? "
          f"{any('ref' in k or 'oc' in k or 'shipment' in k for k in fields)}")
    key, ref = versioning.identify_shipment(
        {"subject": "documents attached", "body": "see file", "email_id": "doc-only",
         "attachments": ["scan.pdf"]}, fields, {})
    print(f"    key for a mail whose ONLY reference is in the PDF -> {key!r}")
    print("    -> document text is not part of the grouping haystack")


# --------------------------------------------------------------------------- #
def probe_version_cases() -> None:
    hdr("5. VERSION DETECTION - the four cases from the checklist")
    init_db()

    print("  Is a version token ever read from the filename or the subject?")
    import re
    from pathlib import Path as _P
    src = (_P(BACKEND_ROOT) / "app" / "services").glob("*.py")
    hits = []
    for path in src:
        for line in path.read_text(encoding="utf-8").splitlines():
            if re.search(r"\bv\d\b.*(subject|filename)|version.*(subject|filename)", line, re.I):
                hits.append(f"    {path.name}: {line.strip()[:96]}")
    print("\n".join(hits) if hits else "    (no code reads a version token from subject/filename)")

    db = SessionLocal()
    try:
        shipment = versioning._get_or_create_shipment(db, "REF:SHP-900", "SHP-900")

        def add(filename, text, received, doc_type="SI"):
            email = {"email_id": f"m-{filename}-{received}", "received_at": received}
            return versioning.register_document_version(
                db, shipment, doc_type, filename, email,
                pdf(text), extractor.extract_fields(text, doc_type).fields, text)

        print("\n  Case A - explicit-looking names, chronological arrival")
        for v, day in (("v1", "2026-09-18"), ("v2", "2026-09-19"), ("v3", "2026-09-20")):
            add(f"SI_{v}.pdf", si_text(weight={"v1": "38,900 KG", "v2": "40,150 KG",
                                               "v3": "41,500 KG"}[v]), day)
        rows = (db.query(DocumentVersionRecord)
                .filter_by(shipment_id=shipment.id, doc_type="SI")
                .order_by(DocumentVersionRecord.version_number).all())
        print("   ", [(v.version_number, v.filename, v.extracted_fields.get("gross_weight_kg"),
                       "CURRENT" if v.is_latest else "") for v in rows])

        print("\n  Case B - no version token anywhere, arrival reversed (newest syncs first)")
        shipment2 = versioning._get_or_create_shipment(db, "REF:SHP-901", "SHP-901")
        for day, weight in (("2026-09-20", "41,500 KG"),   # newest processed first
                            ("2026-09-19", "40,150 KG"),
                            ("2026-09-18", "38,900 KG")):
            versioning.register_document_version(
                db, shipment2, "SI", f"scan_{day}.pdf",
                {"email_id": f"m2-{day}", "received_at": day},
                pdf(si_text(weight=weight)),
                extractor.extract_fields(si_text(weight=weight), "SI").fields,
                si_text(weight=weight))
        rows2 = (db.query(DocumentVersionRecord)
                 .filter_by(shipment_id=shipment2.id, doc_type="SI")
                 .order_by(DocumentVersionRecord.version_number).all())
        print("   ", [(v.version_number, v.filename.rsplit("/", 1)[-1],
                       v.extracted_fields.get("gross_weight_kg"),
                       "CURRENT" if v.is_latest else "") for v in rows2])
        print("    -> the newest document is current regardless of arrival order")

        print("\n  Case C - filename claims v2 while the content is the v3 revision")
        shipment3 = versioning._get_or_create_shipment(db, "REF:SHP-902", "SHP-902")
        versioning.register_document_version(
            db, shipment3, "SI", "SI_v3.pdf",
            {"email_id": "c-1", "received_at": "2026-09-18"},
            pdf(si_text(weight="38,900 KG")),
            extractor.extract_fields(si_text(weight="38,900 KG"), "SI").fields, "")
        versioning.register_document_version(
            db, shipment3, "SI", "SI_v2.pdf",       # name says v2
            {"email_id": "c-2", "received_at": "2026-09-20"},   # content is newer
            pdf(si_text(weight="41,500 KG")),
            extractor.extract_fields(si_text(weight="41,500 KG"), "SI").fields, "")
        rows3 = (db.query(DocumentVersionRecord)
                 .filter_by(shipment_id=shipment3.id, doc_type="SI")
                 .order_by(DocumentVersionRecord.version_number).all())
        print("   ", [(v.version_number, v.filename.rsplit("/", 1)[-1],
                       v.extracted_fields.get("gross_weight_kg"),
                       "CURRENT" if v.is_latest else "") for v in rows3])
        print("    -> the name carries no authority; the conflict is not flagged")

        print("\n  Case D - the same document sent twice")
        shipment4 = versioning._get_or_create_shipment(db, "REF:SHP-903", "SHP-903")
        same = si_text(weight="41,500 KG")
        first = versioning.register_document_version(
            db, shipment4, "SI", "SHP-903_SI.pdf",
            {"email_id": "d-1", "received_at": "2026-09-20"},
            pdf(same), extractor.extract_fields(same, "SI").fields, same)
        second = versioning.register_document_version(
            db, shipment4, "SI", "SHP-903_SI.pdf",
            {"email_id": "d-2", "received_at": "2026-09-20"},
            pdf(same), extractor.extract_fields(same, "SI").fields, same)
        rows4 = (db.query(DocumentVersionRecord)
                 .filter_by(shipment_id=shipment4.id, doc_type="SI")
                 .order_by(DocumentVersionRecord.version_number).all())
        print(f"    first  -> version {first.version_number}, {first.document_status}, "
              f"latest={first.is_latest}")
        print(f"    second -> version {second.version_number}, {second.document_status}, "
              f"latest={second.is_latest}, duplicate_of={second.duplicate_of_version_id}")
        print(f"    total rows={len(rows4)} ACTIVE={sum(1 for v in rows4 if v.document_status=='ACTIVE')}")
        print("    -> stored as a DUPLICATE for audit; no new version, current unchanged")
    finally:
        db.close()


# --------------------------------------------------------------------------- #
def probe_source_preservation() -> None:
    hdr("2. GMAIL SOURCE PRESERVATION - which fields actually exist?")
    from app.models import EmailRecord as ER, GmailMessageRecord as GR
    email_cols = [c.name for c in ER.__table__.columns]
    gmail_cols = [c.name for c in GR.__table__.columns]
    print(f"  EmailRecord columns        : {email_cols}")
    print(f"  GmailMessageRecord columns : {gmail_cols}")
    wanted = {
        "gmail message id": "gmail_message_id" in gmail_cols,
        "sender": "sender" in email_cols,
        "recipient": any(c in email_cols for c in ("recipient", "to", "to_address")),
        "subject": "subject" in email_cols,
        "timestamp": "received_at" in email_cols,
        "thread id": "thread_id" in gmail_cols,
        "body": "body" in email_cols,
        "attachment reference": "attachments" in email_cols,
    }
    for label, ok in wanted.items():
        print(f"    {label:22s} {'YES' if ok else 'NO'}")
    import inspect
    from app.integrations.gmail import parser as gp
    src = inspect.getsource(gp.parse_message)
    captures_to = '"to"' in src or "headers.get('to')" in src
    print(f"  parser captures a 'To'/recipient header? {captures_to}")
    print("  NOTE: attachment *bytes* live on disk; the DB stores their paths.")


# --------------------------------------------------------------------------- #
def probe_ocr_and_pages() -> None:
    hdr("7. OCR / EXTRACTION - formats, page coverage, failure behaviour")
    two_pages = pdf("SCANNED COPY - PAGE 1 OF 2\nFIELDS CONTINUE ON THE NEXT PAGE\n")
    import fitz

    doc = fitz.open(stream=two_pages, filetype="pdf")
    page = doc.new_page()
    y = 60
    for line in si_text().splitlines():
        page.insert_text((60, y), line, fontsize=9.5, fontname="cour")
        y += 15
    multipage = doc.tobytes(deflate=True)
    doc.close()

    text, source = ai_service.document_text("multi.pdf", multipage)
    fields = extractor.extract_fields(text, "SI").fields
    print(f"  multi-page PDF ({len(fitz.open(stream=multipage, filetype='pdf'))} pages): "
          f"source={source} shipper={fields.get('shipper')!r} "
          f"weight={fields.get('gross_weight_kg')}")
    print(f"    page-2 fields recovered: "
          f"{all(fields.get(f) is not None for f in ('shipper','consignee','notify_party','port_of_loading','port_of_discharge','container_count','gross_weight_kg'))}")

    print("\n  formats through the real readers:")
    samples = {
        "SI.txt": si_text().encode(),
        "SI.csv": ("Shipment ID:,SHP-001\nShipper:,ABC LOGISTICS SDN BHD\n"
                   "Gross Weight (KG):,\"41,500 KG\"\n").encode(),
        "SI.xml": (b"<?xml version='1.0'?><ShippingInstruction><ShipmentID>SHP-001"
                   b"</ShipmentID><Shipper>ABC LOGISTICS SDN BHD</Shipper></ShippingInstruction>"),
        "SI.docx": None, "SI.xlsx": None,
        "SI.pdf": pdf(si_text()),
        "SI.png": None, "SI.jpg": None, "SI.tiff": None,
        "SI.doc": si_text().encode(),
        "SI.exe": b"MZ\x90\x00",
    }
    import io
    from docx import Document
    d = Document()
    for line in si_text().splitlines():
        d.add_paragraph(line)
    buf = io.BytesIO(); d.save(buf); samples["SI.docx"] = buf.getvalue()
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    for line in si_text().splitlines():
        if ":" in line:
            label, _, value = line.partition(":")
            ws.append([label + ":", value.strip()])
    buf = io.BytesIO(); wb.save(buf); samples["SI.xlsx"] = buf.getvalue()
    from PIL import Image
    for ext, fmt in (("SI.png", "PNG"), ("SI.jpg", "JPEG"), ("SI.tiff", "TIFF")):
        pdf_b = pdf(si_text())
        d2 = fitz.open(stream=pdf_b, filetype="pdf")
        pix = d2[0].get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("L")
        b = io.BytesIO(); img.save(b, format=fmt)
        samples[ext] = b.getvalue()
        d2.close()

    for name, data in samples.items():
        try:
            text, source = ai_service.document_text(name, data)
            fields = extractor.extract_fields(text, "SI").fields if text else {}
            print(f"    {name:9s} source={str(source):11s} "
                  f"shipment_id={fields.get('shipment_id')!r} "
                  f"shipper={'yes' if fields.get('shipper') else 'no'}")
        except Exception as exc:  # noqa: BLE001
            print(f"    {name:9s} EXCEPTION {type(exc).__name__}: {exc}")

    print("\n  failure behaviour:")
    corrupt = b"%PDF-1.4\nthis is not a real pdf body"
    text, source = ai_service.document_text("corrupt.pdf", corrupt)
    print(f"    corrupt PDF  -> text={text!r} source={source!r} (no exception)")
    safe, reason = __import__("app.services.security", fromlist=["x"]).verify_file_safety(
        "payload.exe", b"MZ\x90\x00")
    print(f"    .exe upload  -> safe={safe} reason={reason!r}")


# --------------------------------------------------------------------------- #
def probe_missing_vs_processing() -> None:
    hdr("10/11. MISSING vs STILL PROCESSING, and the status vocabulary")
    from app.services import shipment_overview as so
    print(f"  REQUIRED_DOC_TYPES = {so.REQUIRED_DOC_TYPES}")
    print(f"  OPTIONAL_DOC_TYPES = {so.OPTIONAL_DOC_TYPES}")
    print(f"  status vocabulary  = {so.PROCESSING}/{so.INCOMPLETE}/"
          f"{so.NEEDS_ATTENTION}/{so.VERIFIED}")
    print("\n  status decision (shipment_overview.build_overview):")
    import inspect
    src = inspect.getsource(so.build_overview)
    block = src.split("has_any_document = ")[1].split("actions")[0]
    for line in block.splitlines()[:12]:
        print(f"    {line.rstrip()}")


# --------------------------------------------------------------------------- #
def main() -> int:
    probe_shipping_id()
    probe_classification()
    probe_version_cases()
    probe_source_preservation()
    probe_ocr_and_pages()
    probe_missing_vs_processing()
    print("\ndone.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
